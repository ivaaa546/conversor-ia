import os
import json
import base64
import io
import requests
from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__, static_folder='.')
CORS(app)

GEMINI_DEFAULT_MODEL = "gemini-2.0-flash"

SYSTEM_PROMPT = """Eres un experto en digitalización de documentos manuscritos y médicos.
Tu objetivo es analizar la imagen del documento cargado (como formularios de consulta externa, registros de pacientes, listas o tablas) y extraer toda la información estructurada con máxima precisión.

REGLAS DE EXTRACCIÓN:
1. Extrae los metadatos del encabezado del documento si existen (ej. Médico, Lugar de Atención, Fecha del documento).
2. Extrae cada fila de la tabla principal en un objeto con los siguientes campos estándar:
   - "no": Número de fila/registro.
   - "nombre": Nombre completo del paciente o sujeto.
   - "cui": CUI, DNI o número de identificación.
   - "fecha_nacimiento": Fecha de nacimiento.
   - "sexo": Sexo/género (M/F si aplica).
   - "telefono": Número telefónico.
   - "domicilio": Dirección o comunidad.
   - "municipio": Municipio o ciudad.
   - "diagnostico": Diagnóstico médico o nota principal.
   - "tratamiento": Tratamiento, medicamento recetado o notas secundarias.

3. Si el formulario tiene columnas adicionales o diferentes, adáptalas o agrégalas como campos descriptivos en el objeto de cada fila.
4. Si la letra es difícil de leer, haz tu mejor esfuerzo basándote en el contexto médico/nombres hispanos y marca palabras dudosas o incompletas sin inventar datos desproporcionados.
5. Si una celda está vacía o tachada con una línea, coloca un string vacío `""` o `"-"`.

RESPONDE ÚNICAMENTE EN FORMATO JSON VÁLIDO CON ESTA ESTRUCTURA EXACTA:
{
  "documento": {
    "titulo": "Título o tipo de documento detectado",
    "medico": "Nombre del médico si existe",
    "lugar": "Lugar de atención si existe",
    "fecha": "Fecha del documento si existe"
  },
  "columnas": ["no", "nombre", "cui", "fecha_nacimiento", "sexo", "telefono", "domicilio", "municipio", "diagnostico", "tratamiento"],
  "filas": [
    {
      "no": "8",
      "nombre": "Sofia Carmelina Pop Xe",
      "cui": "1581 287 241613",
      "fecha_nacimiento": "27/2/87",
      "sexo": "",
      "telefono": "51543814",
      "domicilio": "Sequacil",
      "municipio": "Chirec",
      "diagnostico": "Artralgia",
      "tratamiento": "Aceclofenaco"
    }
  ]
}
"""

@app.route('/')
def index():
    return send_from_directory('.', 'index.html')

@app.route('/<path:path>')
def static_files(path):
    return send_from_directory('.', path)

@app.route('/api/analyze', methods=['POST'])
def analyze_document():
    try:
        # Obtain API Key
        api_key = request.headers.get('X-Gemini-API-Key') or request.form.get('api_key') or os.getenv('GEMINI_API_KEY')
        if not api_key:
            return jsonify({
                "error": "Falta la clave API de Gemini. Ingrésala en la interfaz o configúrala en el archivo .env"
            }), 400

        model = request.form.get('model', GEMINI_DEFAULT_MODEL)
        
        # Check image file upload or base64 payload
        file = request.files.get('file')
        if not file:
            data = request.get_json(silent=True) or {}
            file_b64 = data.get('image_base64')
            mime_type = data.get('mime_type', 'image/jpeg')
            if not file_b64:
                return jsonify({"error": "No se proporcionó ninguna imagen para analizar."}), 400
        else:
            mime_type = file.mimetype or 'image/jpeg'
            file_bytes = file.read()
            file_b64 = base64.b64encode(file_bytes).decode('utf-8')

        # Clean mime_type
        if mime_type == 'application/pdf':
            mime_type = 'application/pdf'
        elif not mime_type.startswith('image/'):
            mime_type = 'image/jpeg'

        # Call Gemini REST API
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
        
        payload = {
            "contents": [
                {
                    "parts": [
                        {"text": SYSTEM_PROMPT},
                        {
                            "inline_data": {
                                "mime_type": mime_type,
                                "data": file_b64
                            }
                        }
                    ]
                }
            ],
            "generationConfig": {
                "temperature": 0.1,
                "response_mime_type": "application/json"
            }
        }

        headers = {"Content-Type": "application/json"}
        response = requests.post(url, headers=headers, json=payload, timeout=60)
        
        if response.status_code != 200:
            err_detail = response.json().get('error', {}).get('message', response.text)
            return jsonify({"error": f"Error de Gemini API ({response.status_code}): {err_detail}"}), response.status_code

        res_data = response.json()
        candidates = res_data.get('candidates', [])
        if not candidates:
            return jsonify({"error": "No se recibió respuesta válida del modelo de IA."}), 500

        text_out = candidates[0].get('content', {}).get('parts', [{}])[0].get('text', '{}')

        # Parse JSON output from Gemini
        try:
            parsed_json = json.loads(text_out)
        except Exception as e:
            # Fallback if raw text returned inside markdown code blocks
            clean_text = text_out.strip()
            if clean_text.startswith("```json"):
                clean_text = clean_text[7:]
            if clean_text.endswith("```"):
                clean_text = clean_text[:-3]
            parsed_json = json.loads(clean_text.strip())

        return jsonify({
            "success": True,
            "data": parsed_json,
            "raw_text": text_out
        })

    except Exception as err:
        return jsonify({"error": f"Error interno en el servidor: {str(err)}"}), 500

@app.route('/api/export-excel', methods=['POST'])
def export_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        req_data = request.get_json()
        filas = req_data.get('filas', [])
        columnas = req_data.get('columnas', [])
        documento = req_data.get('documento', {})

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pacientes Registrados"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name='Calibri', size=16, bold=True, color='1E293B')
        meta_label_font = Font(name='Calibri', size=11, bold=True, color='475569')
        meta_val_font = Font(name='Calibri', size=11, color='0F172A')
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid') # Navy Blue
        
        border_thin = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        # 1. Document Header
        current_row = 1
        doc_title = documento.get('titulo', 'INFORMACIÓN DIARIA DEL PACIENTE')
        ws.cell(row=current_row, column=1, value=doc_title).font = title_font
        current_row += 2

        if documento.get('medico'):
            ws.cell(row=current_row, column=1, value="Médico:").font = meta_label_font
            ws.cell(row=current_row, column=2, value=documento.get('medico')).font = meta_val_font
            current_row += 1

        if documento.get('lugar'):
            ws.cell(row=current_row, column=1, value="Lugar de Atención:").font = meta_label_font
            ws.cell(row=current_row, column=2, value=documento.get('lugar')).font = meta_val_font
            current_row += 1

        if documento.get('fecha'):
            ws.cell(row=current_row, column=1, value="Fecha:").font = meta_label_font
            ws.cell(row=current_row, column=2, value=documento.get('fecha')).font = meta_val_font
            current_row += 1

        current_row += 1 # Empty spacing row

        # Mapping column keys to friendly labels
        headers_map = {
            "no": "No.",
            "nombre": "NOMBRE DEL PACIENTE",
            "cui": "CUI / DPI",
            "fecha_nacimiento": "FECHA DE NACIMIENTO",
            "sexo": "SEXO",
            "telefono": "TELÉFONO",
            "domicilio": "DOMICILIO",
            "municipio": "MUNICIPIO",
            "diagnostico": "DIAGNÓSTICO",
            "tratamiento": "TRATAMIENTO"
        }

        if not columnas and filas:
            columnas = list(filas[0].keys())

        # 2. Table Headers
        for col_idx, col_key in enumerate(columnas, start=1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = headers_map.get(col_key.lower(), col_key.upper())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border
        
        ws.row_dimensions[current_row].height = 26
        current_row += 1

        # 3. Table Rows
        zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        
        for r_idx, fila in enumerate(filas):
            for c_idx, col_key in enumerate(columnas, start=1):
                val = fila.get(col_key, '')
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.border = cell_border
                cell.alignment = Alignment(vertical="center")
                if r_idx % 2 == 1:
                    cell.fill = zebra_fill
            current_row += 1

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="pacientes_digitalizados.xlsx"
        )

    except Exception as e:
        return jsonify({"error": f"Error al generar archivo Excel: {str(e)}"}), 500

if __name__ == '__main__':
    port = int(os.getenv('PORT', 5050))
    print(f"🚀 Servidor ejecutándose en http://localhost:{port}")
    app.run(host='0.0.0.0', port=port, debug=False)


